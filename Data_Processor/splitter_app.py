# ================================================================
# MASTER EXCEL
# CREATE TRANSECT-WISE AND STATION-WISE SHEETS  (+ 0 m SHEET)
# Python / Streamlit version of Splitter.R
# ================================================================
#
# WHAT THIS APP DOES
#
# 1. You upload your MASTER Excel file in the browser.
# 2. Reads the first sheet of the Excel file (you may pick another).
# 3. Automatically finds the column containing station names
#    (T1 S1, T1 S2, T2 S1 ...).
# 4. Creates TRANSECT sheets   : Transect_T1, Transect_T2, ...
# 5. Creates STATION sheets    : Station_S1, Station_S2, ...
# 6. NEW: creates one extra sheet "Depth_0m" containing every
#    0 m record from all transects and stations.
# 7. Every original parameter remains with its station record.
# 8. Original master sheet is retained.
# 9. The workbook with the new sheets is offered for download.
#
# RUN:
#     pip install streamlit pandas openpyxl
#     streamlit run splitter_app.py
# ================================================================

import io
import re
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ZERO_SHEET_NAME = "Depth_0m"

# ================================================================
# PAGE SETUP
# ================================================================

st.set_page_config(
    page_title="Transect & Station Splitter",
    page_icon="🌊",
    layout="wide",
)

st.markdown(
    """
    <style>
    .block-container {padding-top: 2rem; max-width: 1200px;}
    .hero {
        background: linear-gradient(120deg, #0b3c5d 0%, #1d7a99 60%, #3fb8af 100%);
        padding: 1.6rem 2rem; border-radius: 16px; color: white;
        margin-bottom: 1.4rem;
    }
    .hero h1 {margin: 0; font-size: 2rem; color: white;}
    .hero p {margin: .3rem 0 0 0; opacity: .92; font-size: 1.02rem;}
    div[data-testid="stMetric"] {
        background: #f3f8fb; border: 1px solid #d7e6ee;
        border-radius: 12px; padding: .8rem 1rem;
    }
    div[data-testid="stMetricLabel"] p {color: #33566a;}
    div[data-testid="stMetricValue"] {color: #0b3c5d;}
    .stDownloadButton button, .stButton button[kind="primary"] {
        border-radius: 10px; font-weight: 600;
    }
    </style>
    <div class="hero">
        <h1>🌊 Transect &amp; Station Splitter</h1>
        <p>Master Excel &rarr; Transect sheets, Station sheets and a 0&nbsp;m sheet</p>
    </div>
    """,
    unsafe_allow_html=True,
)


# ================================================================
# HELPER FUNCTIONS
# ================================================================

def clean_column_name(name):
    """Trim, lowercase, replace _ and - by space, squeeze spaces."""
    n = str(name).strip().lower()
    n = n.replace("_", " ")
    n = n.replace("-", " ")
    n = re.sub(r"\s+", " ", n)
    return n.strip()


def find_station_column(columns):
    """SECTION 5 : find station-name column automatically."""
    clean_names = [clean_column_name(c) for c in columns]

    possible_station_names = [
        "station name",
        "station",
        "stationname",
        "station id",
        "stationid",
    ]

    # exact matching column
    idx = [i for i, n in enumerate(clean_names) if n in possible_station_names]

    # otherwise any column containing "station"
    if len(idx) == 0:
        idx = [i for i, n in enumerate(clean_names) if "station" in n]

    if len(idx) == 0:
        return None
    return columns[idx[0]]


def find_depth_column(columns):
    """NEW : find the depth column (used for the 0 m sheet)."""
    clean_names = [clean_column_name(c) for c in columns]

    exact = ["depth", "depth m", "depth (m)", "depthm", "depth in m", "sampling depth"]
    idx = [i for i, n in enumerate(clean_names) if n in exact]

    if len(idx) == 0:
        idx = [i for i, n in enumerate(clean_names) if "depth" in n]

    if len(idx) == 0:
        return None
    return columns[idx[0]]


def is_zero_depth(value):
    """True for 0, 0.0, '0', '0m', '0 m', '0 M', '0.0 m' ..."""
    if value is None:
        return False
    try:
        if pd.isna(value):
            return False
    except (TypeError, ValueError):
        return False
    s = str(value).strip().lower().replace(" ", "")
    s = re.sub(r"(meters|meter|metres|metre|m)$", "", s)
    try:
        return float(s.replace(",", ".")) == 0
    except ValueError:
        return False


def to_cell(v):
    """Convert pandas / numpy values into plain Excel-writable values."""
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item"):
        return v.item()
    return v


def write_sheet(wb, sheet_name, data):
    """Add worksheet, write data with filter, bold centred header,
    frozen first row and automatic column widths."""
    ws = wb.create_sheet(title=sheet_name)

    ws.append([str(c) for c in data.columns])
    for row in data.itertuples(index=False, name=None):
        ws.append([to_cell(v) for v in row])

    ncol = len(data.columns)

    # Format header
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Filter + freeze first row
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{max(len(data) + 1, 1)}"
    ws.freeze_panes = "A2"

    # Automatic column widths + date formatting
    for i in range(1, ncol + 1):
        col_values = [ws.cell(row=r, column=i).value for r in range(1, min(ws.max_row, 2001) + 1)]
        longest = max(len(str(v)) if v is not None else 0 for v in col_values)
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 8), 60)

        dt_cells = [
            ws.cell(row=r, column=i)
            for r in range(2, ws.max_row + 1)
            if hasattr(ws.cell(row=r, column=i).value, "hour")
        ]
        if dt_cells and all(
            c.value.hour == 0 and c.value.minute == 0 and c.value.second == 0
            for c in dt_cells
            if hasattr(c.value, "hour")
        ):
            for c in dt_cells:
                c.number_format = "yyyy-mm-dd"


def num_key(text, prefix):
    return int(re.sub(f"^{prefix}", "", text))


# ================================================================
# MAIN PROCESSING
# ================================================================

def run_split(file_bytes, sheet_name, station_column, depth_column):
    log = []
    log_add = log.append

    # ------------------------------------------------------------
    # SECTION 3 : READ MASTER EXCEL
    # ------------------------------------------------------------
    log_add("Reading master Excel file...")
    master_data = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_name)

    if len(master_data) == 0:
        raise ValueError("ERROR: The selected Excel sheet contains no data.")
    log_add("Master Excel file successfully read.\n")

    # ------------------------------------------------------------
    # SECTION 4 : DISPLAY COLUMN NAMES
    # ------------------------------------------------------------
    log_add("====================================================")
    log_add(" COLUMNS FOUND IN MASTER FILE")
    log_add("====================================================\n")
    log_add(str(list(master_data.columns)))
    log_add("")
    log_add(f"Station-name column detected as: {station_column}\n")

    # ------------------------------------------------------------
    # SECTION 6 : CLEAN STATION-NAME COLUMN
    # ------------------------------------------------------------
    master_data[station_column] = master_data[station_column].apply(
        lambda v: "" if (v is None or (not isinstance(v, str) and pd.isna(v))) else str(v).strip()
    )
    master_data = master_data[master_data[station_column] != ""].copy()

    # ------------------------------------------------------------
    # SECTION 7 : EXTRACT TRANSECT AND STATION
    # ------------------------------------------------------------
    master_data["Transect_ID"] = master_data[station_column].apply(
        lambda s: re.sub(r"^([Tt][0-9]+)\s+.*$", r"\1", s)
    )
    master_data["Station_ID"] = master_data[station_column].apply(
        lambda s: re.sub(r"^[Tt][0-9]+\s+([Ss][0-9]+).*$", r"\1", s)
    )
    master_data["Transect_ID"] = master_data["Transect_ID"].str.upper()
    master_data["Station_ID"] = master_data["Station_ID"].str.upper()

    # ------------------------------------------------------------
    # SECTION 8 : CHECK STATION-NAME FORMAT
    # ------------------------------------------------------------
    valid_format = master_data[station_column].apply(
        lambda s: re.fullmatch(r"[Tt][0-9]+\s+[Ss][0-9]+", s) is not None
    )
    invalid_rows = master_data[~valid_format]
    invalid_names = list(pd.unique(invalid_rows[station_column]))

    if len(invalid_rows) > 0:
        log_add("")
        log_add("WARNING: Some station names do not follow the expected format.")
        log_add("Expected format: T1 S1, T1 S2, T2 S1, etc.\n")
        log_add("Station names requiring checking:\n")
        log_add(str(invalid_names))
        log_add("")

    # ------------------------------------------------------------
    # SECTION 9 : FIND ALL TRANSECTS
    # ------------------------------------------------------------
    transects = [t for t in pd.unique(master_data["Transect_ID"]) if re.fullmatch(r"T[0-9]+", t)]
    transects = sorted(transects, key=lambda t: num_key(t, "T"))

    # ------------------------------------------------------------
    # SECTION 10 : FIND ALL STATIONS
    # ------------------------------------------------------------
    stations = [s for s in pd.unique(master_data["Station_ID"]) if re.fullmatch(r"S[0-9]+", s)]
    stations = sorted(stations, key=lambda s: num_key(s, "S"))

    # ------------------------------------------------------------
    # SECTION 11 : SHOW DETECTED STRUCTURE
    # ------------------------------------------------------------
    log_add("")
    log_add("====================================================")
    log_add(" DETECTED DATA STRUCTURE")
    log_add("====================================================\n")
    log_add(f"Total rows:  {len(master_data)}")
    log_add(f"Total original columns:  {master_data.shape[1] - 2}")
    log_add(f"Transects detected:  {', '.join(transects)}")
    log_add(f"Stations detected:  {', '.join(stations)}\n")

    # ------------------------------------------------------------
    # SECTION 12 : SHOW TRANSECT / STATION COMBINATIONS
    # ------------------------------------------------------------
    log_add("Transect / Station structure:\n")
    structure = []
    for transect in transects:
        temp = master_data[master_data["Transect_ID"] == transect]
        temp_stations = [s for s in pd.unique(temp["Station_ID"]) if re.fullmatch(r"S[0-9]+", s)]
        temp_stations = sorted(temp_stations, key=lambda s: num_key(s, "S"))
        log_add(f"{transect}  ->  {', '.join(temp_stations)}")
        structure.append({"Transect": transect, "Stations": ", ".join(temp_stations)})
    log_add("")

    # ------------------------------------------------------------
    # SECTION 13 : LOAD ORIGINAL WORKBOOK
    # ------------------------------------------------------------
    wb = load_workbook(io.BytesIO(file_bytes))

    # ------------------------------------------------------------
    # SECTION 14 : REMOVE PREVIOUSLY GENERATED SHEETS
    # ------------------------------------------------------------
    generated_sheets = [
        s for s in wb.sheetnames
        if re.fullmatch(r"(Transect_T[0-9]+|Station_S[0-9]+|" + ZERO_SHEET_NAME + r")", s)
    ]
    if len(generated_sheets) > 0:
        log_add("Removing previously generated sheets...\n")
        for old_sheet in generated_sheets:
            del wb[old_sheet]
            log_add(f"Removed: {old_sheet}")

    # ------------------------------------------------------------
    # SECTION 16 : CREATE TRANSECT SHEETS
    # ------------------------------------------------------------
    log_add("")
    log_add("====================================================")
    log_add(" CREATING TRANSECT SHEETS")
    log_add("====================================================\n")

    for transect in transects:
        transect_data = master_data[master_data["Transect_ID"] == transect]
        transect_data = transect_data.drop(columns=["Transect_ID", "Station_ID"])
        sheet = f"Transect_{transect}"
        write_sheet(wb, sheet, transect_data)
        log_add(f"Created: {sheet} | Rows: {len(transect_data)}")

    # ------------------------------------------------------------
    # SECTION 17 : CREATE STATION SHEETS
    # ------------------------------------------------------------
    log_add("")
    log_add("====================================================")
    log_add(" CREATING STATION SHEETS")
    log_add("====================================================\n")

    for station in stations:
        station_data = master_data[master_data["Station_ID"] == station]
        station_data = station_data.drop(columns=["Transect_ID", "Station_ID"])
        sheet = f"Station_{station}"
        write_sheet(wb, sheet, station_data)
        log_add(f"Created: {sheet} | Rows: {len(station_data)}")

    # ------------------------------------------------------------
    # NEW SECTION : CREATE 0 m SHEET
    # All 0 m records from all transects and stations.
    # ------------------------------------------------------------
    zero_data = None
    if depth_column is not None:
        log_add("")
        log_add("====================================================")
        log_add(" CREATING 0 m SHEET")
        log_add("====================================================\n")
        log_add(f"Depth column detected as: {depth_column}\n")

        mask = master_data[depth_column].apply(is_zero_depth)
        zero_data = master_data[mask].drop(columns=["Transect_ID", "Station_ID"])
        write_sheet(wb, ZERO_SHEET_NAME, zero_data)
        log_add(f"Created: {ZERO_SHEET_NAME} | Rows: {len(zero_data)}")
        if len(zero_data) == 0:
            log_add("WARNING: No 0 m records were found in the depth column.")
    else:
        log_add("\nWARNING: No depth column selected - 0 m sheet was not created.")

    # ------------------------------------------------------------
    # SECTION 18 : SAVE WORKBOOK (to memory, offered for download)
    # ------------------------------------------------------------
    out = io.BytesIO()
    wb.save(out)

    # ------------------------------------------------------------
    # SECTION 19 : FINAL SUMMARY
    # ------------------------------------------------------------
    new_sheets = len(transects) + len(stations) + (1 if zero_data is not None else 0)

    log_add("")
    log_add("====================================================")
    log_add(" PROCESS COMPLETED SUCCESSFULLY")
    log_add("====================================================\n")
    log_add("Original master sheet: RETAINED\n")
    log_add("Transect sheets created:")
    for t in transects:
        log_add(f"  - Transect_{t}")
    log_add("\nStation sheets created:")
    for s in stations:
        log_add(f"  - Station_{s}")
    if zero_data is not None:
        log_add(f"\n0 m sheet created:\n  - {ZERO_SHEET_NAME}")
    log_add("")
    log_add(f"Number of transects: {len(transects)}")
    log_add(f"Number of stations: {len(stations)}")
    log_add(f"Number of new sheets: {new_sheets}\n")
    log_add("All original parameters have been retained.")
    log_add("Original row order has been preserved.")

    return {
        "bytes": out.getvalue(),
        "log": "\n".join(log),
        "transects": transects,
        "stations": stations,
        "structure": structure,
        "invalid_names": invalid_names,
        "rows": len(master_data),
        "zero_data": zero_data,
        "new_sheets": new_sheets,
    }


# ================================================================
# USER INTERFACE
# ================================================================

st.subheader("1 · Select MASTER Excel file")
uploaded = st.file_uploader(
    "Upload your MASTER Excel file",
    type=["xlsx", "xlsm"],
    help="The original master sheet is kept; the new sheets are added to a copy you can download.",
)

if uploaded is None:
    st.info("Upload the MASTER Excel file to begin.")
    st.stop()

file_bytes = uploaded.getvalue()

try:
    sheet_names = load_workbook(io.BytesIO(file_bytes), read_only=True).sheetnames
except Exception as e:
    st.error(f"Could not open the Excel file: {e}")
    st.stop()

c1, c2, c3 = st.columns(3)
with c1:
    sheet_choice = st.selectbox("Master sheet (default = first sheet)", sheet_names, index=0)

try:
    preview = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet_choice)
except Exception as e:
    st.error(f"Could not read the sheet: {e}")
    st.stop()

if len(preview) == 0:
    st.error("ERROR: The selected Excel sheet contains no data.")
    st.stop()

columns = list(preview.columns)
auto_station = find_station_column(columns)
auto_depth = find_depth_column(columns)

if auto_station is None:
    st.error(
        "ERROR: Could not identify the station-name column.\n\n"
        "Columns found in the Excel file are:\n\n"
        + "\n".join(f"- {c}" for c in columns)
        + "\n\nPlease check which column contains values such as 'T1 S1', 'T1 S2', 'T2 S1', etc."
    )
    st.stop()

with c2:
    station_column = st.selectbox(
        "Station-name column (auto-detected)", columns, index=columns.index(auto_station)
    )
with c3:
    depth_options = ["(none - skip 0 m sheet)"] + columns
    depth_choice = st.selectbox(
        "Depth column for 0 m sheet (auto-detected)",
        depth_options,
        index=(depth_options.index(auto_depth) if auto_depth is not None else 0),
    )
depth_column = None if depth_choice == depth_options[0] else depth_choice

with st.expander("Preview of master sheet", expanded=False):
    st.dataframe(preview.head(50), use_container_width=True)

st.subheader("2 · Create sheets")
if st.button("Create Transect, Station & 0 m sheets", type="primary"):
    try:
        with st.spinner("Processing..."):
            st.session_state["result"] = run_split(
                file_bytes, sheet_choice, station_column, depth_column
            )
            st.session_state["result_name"] = uploaded.name
    except Exception as e:
        st.session_state.pop("result", None)
        st.error(str(e))

res = st.session_state.get("result")
if res is not None and st.session_state.get("result_name") == uploaded.name:
    st.success("PROCESS COMPLETED SUCCESSFULLY – original master sheet retained.")

    if res["invalid_names"]:
        st.warning(
            "Some station names do not follow the expected format (T1 S1, T1 S2, T2 S1, etc.):\n\n"
            + ", ".join(f"`{n}`" for n in res["invalid_names"])
        )

    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric("Rows", res["rows"])
    m2.metric("Transects", len(res["transects"]))
    m3.metric("Stations", len(res["stations"]))
    m4.metric("0 m rows", 0 if res["zero_data"] is None else len(res["zero_data"]))
    m5.metric("New sheets", res["new_sheets"])

    stem = Path(uploaded.name).stem
    st.download_button(
        "⬇️  Download Excel file with new sheets",
        data=res["bytes"],
        file_name=f"{stem}_Split.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    tab1, tab2, tab3 = st.tabs(["Transect / Station structure", "0 m sheet preview", "Processing log"])
    with tab1:
        st.dataframe(pd.DataFrame(res["structure"]), use_container_width=True, hide_index=True)
    with tab2:
        if res["zero_data"] is None:
            st.info("No depth column was selected, so the 0 m sheet was not created.")
        else:
            st.caption(f"Sheet name: {ZERO_SHEET_NAME}")
            st.dataframe(res["zero_data"], use_container_width=True)
    with tab3:
        st.code(res["log"], language="text")
