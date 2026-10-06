# ============================================================
# CTD .ASC FILE TO EXCEL DEPTH EXTRACTOR  (Streamlit app)
# ============================================================
#
# Python / Streamlit version of DEPTH-EXTRACTOR.R
#
# Reads a CTD .asc (or .ask) file, identifies the parameters
# automatically, separates downcast and upcast, extracts the
# observations closest to the selected standard depths and
# writes an Excel workbook with TWO sheets:
#
#   Sheet 1 "Selected_Depths" : same output as the R code
#   Sheet 2 "Complete_Data"   : the complete CTD file converted
#                               to Excel (every data row, every
#                               parameter, original file order)
#
# Run with:   streamlit run ctd_depth_extractor_app.py
# ============================================================

import io
import re

import numpy as np
import pandas as pd
import streamlit as st

DEFAULT_STANDARD_DEPTHS = [5, 10, 20, 30, 50, 75, 100, 200, 500, 750, 1000]

SHEET_SELECTED = "Selected_Depths"
SHEET_COMPLETE = "Complete_Data"

# Same patterns as the R code
HEADER_PATTERN = re.compile(r"(^|\s)DepSM(\s|$)", re.IGNORECASE)
NUMERIC_START_PATTERN = re.compile(
    r"^\s*[-+]?([0-9]*\.?[0-9]+|[0-9]+\.?[0-9]*)([eE][-+]?[0-9]+)?"
)


# ============================================================
# HELPERS
# ============================================================

def make_unique(names, sep="_"):
    """Equivalent of R's make.unique(): a, a_1, a_2 ..."""
    seen = {}
    result = []
    used = set(names)
    for name in names:
        if name not in seen:
            seen[name] = 0
            result.append(name)
        else:
            seen[name] += 1
            candidate = f"{name}{sep}{seen[name]}"
            while candidate in used:
                seen[name] += 1
                candidate = f"{name}{sep}{seen[name]}"
            used.add(candidate)
            result.append(candidate)
    return result


def decode_bytes(raw: bytes) -> str:
    """Decode the uploaded file, tolerating odd encodings."""
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


# ============================================================
# 1. READ AND PARSE THE CTD FILE
# ============================================================

def read_ctd_file(raw: bytes):
    """
    Returns (raw_data, column_names, depth_column, info)

    raw_data has one column per CTD parameter plus
    'OriginalFileRow' (1..n, original order).
    """
    text = decode_bytes(raw)
    all_lines = text.splitlines()

    # Remove empty lines
    all_lines = [ln for ln in all_lines if ln.strip()]

    if len(all_lines) < 2:
        raise ValueError("The ASC file does not contain enough data.")

    # ---- Find the header (line containing DepSM, first 100 lines)
    header_index = None
    for i, line in enumerate(all_lines[:100]):
        if HEADER_PATTERN.search(line):
            header_index = i
            break

    if header_index is None:
        raise ValueError(
            "Could not find the CTD header containing 'DepSM'. "
            "Please check the ASC file format."
        )

    header_line = all_lines[header_index]

    # ---- Read the header (tabs or spaces)
    column_names = [c for c in header_line.strip().split() if c]
    column_names = make_unique(column_names, sep="_")

    # ---- Data lines: keep only lines starting with a number
    data_lines = [
        ln for ln in all_lines[header_index + 1:]
        if NUMERIC_START_PATTERN.match(ln)
    ]

    if not data_lines:
        raise ValueError("No numerical CTD data rows were found.")

    # ---- Split into columns (any whitespace), pad ragged rows
    split_rows = [ln.split() for ln in data_lines]
    n_data_cols = max(len(r) for r in split_rows)
    split_rows = [r + [None] * (n_data_cols - len(r)) for r in split_rows]
    raw_data = pd.DataFrame(split_rows)

    # ---- Check column count
    warning = None
    if raw_data.shape[1] < len(column_names):
        raise ValueError(
            "The CTD data contain fewer columns than the header.\n"
            f"Header columns: {len(column_names)}\n"
            f"Data columns: {raw_data.shape[1]}"
        )

    if raw_data.shape[1] > len(column_names):
        warning = (
            f"The data contain {raw_data.shape[1]} columns, while the "
            f"header contains {len(column_names)} columns. "
            "Extra columns were ignored."
        )
        raw_data = raw_data.iloc[:, : len(column_names)]

    raw_data.columns = column_names

    # ---- Find the depth column
    depth_column = None
    for name in column_names:
        if name.strip().lower() == "depsm":
            depth_column = name
            break
    if depth_column is None:
        raise ValueError("The 'DepSM' depth column could not be found.")

    # ---- Convert everything to numeric (handles 7.2180e+02)
    for col in raw_data.columns:
        raw_data[col] = pd.to_numeric(raw_data[col], errors="coerce")

    # ---- Remove rows without a valid depth
    raw_data = raw_data[raw_data[depth_column].notna()].reset_index(drop=True)

    # ---- Add original file row number
    raw_data["OriginalFileRow"] = np.arange(1, len(raw_data) + 1)

    if raw_data.empty:
        raise ValueError("No valid depth observations were found.")

    info = {
        "n_parameters": len(column_names),
        "header_line_number": header_index + 1,
        "warning": warning,
    }
    return raw_data, column_names, depth_column, info


# ============================================================
# 2. SPLIT INTO DOWNCAST / UPCAST
# ============================================================

def split_casts(raw_data, depth_column):
    """
    Everything up to and including the deepest point = DOWNCAST
    The deepest point onward                           = UPCAST
    (the deepest observation belongs to both)
    """
    depth = raw_data[depth_column].to_numpy()
    max_index = int(np.argmax(depth))  # first occurrence, like which.max

    downcast = raw_data.iloc[: max_index + 1]

    if max_index < len(raw_data) - 1:
        upcast = raw_data.iloc[max_index:]
    else:
        upcast = raw_data.iloc[0:0]

    return downcast, upcast


# ============================================================
# 3. SELECT THE CLOSEST OBSERVATIONS
# ============================================================

def select_closest_depths(data, cast_name, standard_depths, depth_column):
    if len(data) == 0:
        return None

    d = data[depth_column].to_numpy()
    cast_min = float(np.min(d))
    cast_max = float(np.max(d))

    available = [s for s in standard_depths if cast_min <= s <= cast_max]

    # Actual minimum, standard depths available, actual maximum
    targets = list(dict.fromkeys([cast_min, *available, cast_max]))

    rows = []
    for target in targets:
        diff = np.abs(d - target)
        closest = int(np.argmin(diff))  # first minimum, like which.min

        row = data.iloc[[closest]].copy()
        row["Cast"] = cast_name
        row["TargetDepth_m"] = target
        row["DepthDifference_m"] = diff[closest]
        rows.append(row)

    return pd.concat(rows, ignore_index=True)


def build_selected_sheet(raw_data, column_names, depth_column, standard_depths):
    downcast, upcast = split_casts(raw_data, depth_column)

    down_sel = select_closest_depths(downcast, "Downcast", standard_depths, depth_column)
    up_sel = select_closest_depths(upcast, "Upcast", standard_depths, depth_column)

    parts = [p for p in (down_sel, up_sel) if p is not None]
    selected = pd.concat(parts, ignore_index=True)

    # Restore original file order (stable, no depth sorting)
    selected = selected.sort_values("OriginalFileRow", kind="stable").reset_index(drop=True)

    # Reorder columns
    ctd_columns = [c for c in column_names if c in selected.columns]
    order = ["Cast", "TargetDepth_m", depth_column, "DepthDifference_m",
             *ctd_columns, "OriginalFileRow"]
    order = [c for c in dict.fromkeys(order) if c in selected.columns]

    return selected[order], len(downcast), len(upcast)


# ============================================================
# 4. SHEET 2: COMPLETE FILE
# ============================================================

def build_complete_sheet(raw_data, column_names):
    """Complete CTD file: every data row, every parameter, original order."""
    cols = [*column_names, "OriginalFileRow"]
    return raw_data[cols].copy()


# ============================================================
# 5. WRITE EXCEL (TWO SHEETS)
# ============================================================

def _format_sheet(ws):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", start_color="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for idx, cell in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(12, len(str(cell.value)) + 3)


def write_excel(selected_data, complete_data) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        selected_data.to_excel(writer, sheet_name=SHEET_SELECTED, index=False)
        complete_data.to_excel(writer, sheet_name=SHEET_COMPLETE, index=False)
        _format_sheet(writer.sheets[SHEET_SELECTED])
        _format_sheet(writer.sheets[SHEET_COMPLETE])
    return buffer.getvalue()


def parse_depth_list(text):
    values = []
    for token in re.split(r"[,\s;]+", text.strip()):
        if token:
            values.append(float(token))
    if not values:
        raise ValueError("Please enter at least one standard depth.")
    return sorted(set(values))


# ============================================================
# STREAMLIT APP
# ============================================================

def main():
    st.set_page_config(page_title="CTD Depth Extractor", page_icon="🌊", layout="wide")

    st.title("🌊 CTD .ASC to Excel Depth Extractor")
    st.write(
        "Upload a CTD `.asc` (or `.ask`) file. The app detects the parameters, "
        "separates downcast and upcast, picks the observations closest to the "
        "standard depths, and creates an Excel file with two sheets:"
    )
    st.markdown(
        f"- **{SHEET_SELECTED}**: selected depths (downcast and upcast)\n"
        f"- **{SHEET_COMPLETE}**: the complete CTD file converted to Excel"
    )

    with st.sidebar:
        st.header("Settings")
        depth_text = st.text_area(
            "Standard depths (m)",
            value=", ".join(str(d) for d in DEFAULT_STANDARD_DEPTHS),
            help="Separate values with commas or spaces. The actual minimum and "
                 "maximum depths of each cast are always included.",
        )

    uploaded = st.file_uploader(
        "Select the CTD file", type=["asc", "ask", "txt", "cnv"], accept_multiple_files=False
    )

    if uploaded is None:
        st.info("Upload a file to begin.")
        return

    try:
        standard_depths = parse_depth_list(depth_text)
    except ValueError as e:
        st.error(f"Invalid standard depth list: {e}")
        return

    try:
        raw_data, column_names, depth_column, info = read_ctd_file(uploaded.getvalue())
        selected_data, n_down, n_up = build_selected_sheet(
            raw_data, column_names, depth_column, standard_depths
        )
        complete_data = build_complete_sheet(raw_data, column_names)
        excel_bytes = write_excel(selected_data, complete_data)
    except Exception as e:
        st.error(f"ERROR: {e}")
        return

    if info["warning"]:
        st.warning(info["warning"])

    depth = raw_data[depth_column]
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Parameters", info["n_parameters"])
    c2.metric("Min depth (m)", f"{depth.min():.3f}")
    c3.metric("Max depth (m)", f"{depth.max():.3f}")
    c4.metric("Downcast rows", n_down)
    c5.metric("Upcast rows", n_up)

    out_name = uploaded.name.rsplit(".", 1)[0] + "_selected_depths.xlsx"

    st.download_button(
        "⬇️ Download Excel file",
        data=excel_bytes,
        file_name=out_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )

    tab1, tab2 = st.tabs([
        f"{SHEET_SELECTED} ({len(selected_data)} rows)",
        f"{SHEET_COMPLETE} ({len(complete_data)} rows)",
    ])
    with tab1:
        st.dataframe(selected_data, use_container_width=True)
    with tab2:
        st.dataframe(complete_data, use_container_width=True)


if __name__ == "__main__":
    main()
