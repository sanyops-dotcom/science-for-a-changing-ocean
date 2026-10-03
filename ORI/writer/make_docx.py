"""ORI Phase 6 - make_docx.py: convert the shipped ori_result.md into a Word (.docx) file.

Run after main.py:  python -m writer.make_docx
Pure Python (python-docx) - no external program required, unlike the earlier pandoc-based
version. Install once with: pip install python-docx
Only converts a result that already shipped (zero verification violations) - see run_all.py.
"""
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    print("Missing dependency. Install it with:\n    pip install python-docx", file=sys.stderr)
    raise

import config

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _add_runs(paragraph, text: str) -> None:
    """Split on **bold** markers and add each piece as a run, bold or not."""
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _bullet_level(line: str) -> int:
    """Count leading 2-space groups before a '- ' bullet marker (0, 1, 2, ...)."""
    stripped = line.lstrip(" ")
    indent = len(line) - len(stripped)
    return indent // 2


def md_to_docx(md_text: str, title: str = None) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for raw_line in md_text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        stripped = line.strip()

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("- ") or stripped.startswith("-"):
            level = min(_bullet_level(line), 3)
            text = stripped[1:].strip() if stripped.startswith("- ") else stripped[1:].strip()
            style_name = "List Bullet" if level == 0 else f"List Bullet {min(level + 1, 3)}"
            try:
                p = doc.add_paragraph(style=style_name)
            except KeyError:
                p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, text)
        else:
            p = doc.add_paragraph()
            _add_runs(p, stripped)
    return doc


def make_docx(out_dir: Path = None) -> Path:
    out_dir = Path(out_dir or config.OUTPUT_DIR) / "writer"
    src = out_dir / "ori_result.md"
    if not src.exists():
        raise FileNotFoundError(f"{src} not found - run main.py first, and check it shipped (no violations).")
    dst = out_dir / "ori_result.docx"
    doc = md_to_docx(src.read_text(encoding="utf-8"))
    doc.save(str(dst))
    return dst


if __name__ == "__main__":
    path = make_docx()
    print(f"Wrote {path}")
